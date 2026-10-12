"""Builds presentation/AlertLens.pptx from presentation/figs. Run make_figures.py first."""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
FIGS = HERE / "figs"
NAVY, RED, BLUE, GRAY = RGBColor(0x14, 0x21, 0x3D), RGBColor(0xD7, 0x26, 0x3D), RGBColor(0x3A, 0x7D, 0xC9), RGBColor(0x5B, 0x67, 0x75)
LIGHT, WHITE, ORANGE = RGBColor(0xF3, 0xF5, 0xF8), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xF4, 0xA2, 0x61)
FONT = "Helvetica Neue"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
page = 0


def text(slide, x, y, w, h, s, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for k, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = color
    return tb


def bullets(slide, x, y, w, h, items, size=18, color=NAVY, gap=10):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for k, item in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, _, rest = item.partition("|")
        r = p.add_run()
        r.text = "■  "
        r.font.size, r.font.color.rgb, r.font.name = Pt(size - 6), RED, FONT
        r = p.add_run()
        r.text = head
        r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bool(rest), FONT, color
        if rest:
            r = p.add_run()
            r.text = " " + rest
            r.font.size, r.font.name, r.font.color.rgb = Pt(size), FONT, color
    return tb


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def picture(slide, name, x, y, max_w, max_h):
    iw, ih = Image.open(FIGS / name).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(str(FIGS / name), x + (max_w - w) // 2, y + (max_h - h) // 2, w, h)


def stat(slide, x, y, w, big, small, color=RED):
    text(slide, x, y, w, Inches(0.9), big, size=40, color=color, bold=True)
    text(slide, x, y + Inches(0.85), w, Inches(0.9), small, size=14, color=GRAY)


def new_slide(title, kicker=None, notes=""):
    global page
    page += 1
    s = prs.slides.add_slide(BLANK)
    if kicker:
        text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.4), kicker.upper(), size=13, color=RED, bold=True)
    text(s, Inches(0.6), Inches(0.65), Inches(12.2), Inches(0.9), title, size=30, bold=True)
    box(s, Inches(0.6), Inches(1.45), Inches(1.1), Inches(0.07), RED, shape=MSO_SHAPE.RECTANGLE)
    text(s, Inches(0.6), H - Inches(0.5), Inches(8), Inches(0.35), "AlertLens  ·  AIC × CrowdStrike  ·  Track 2: Detect",
         size=11, color=GRAY)
    text(s, W - Inches(1.4), H - Inches(0.5), Inches(0.8), Inches(0.35), str(page), size=11, color=GRAY,
         align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes.strip()
    return s


def callout(slide, x, y, w, h, s, size=17):
    box(slide, x, y, w, h, LIGHT)
    box(slide, x, y, Inches(0.09), h, RED, shape=MSO_SHAPE.RECTANGLE)
    text(slide, x + Inches(0.25), y, w - Inches(0.35), h, s, size=size, anchor=MSO_ANCHOR.MIDDLE)


# 1 ─ title
page += 1
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, H, NAVY, shape=MSO_SHAPE.RECTANGLE)
box(s, Inches(0.8), Inches(3.55), Inches(1.4), Inches(0.09), RED, shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(0.8), Inches(1.6), Inches(11), Inches(1.2), "AlertLens", size=66, color=WHITE, bold=True)
text(s, Inches(0.8), Inches(2.65), Inches(11.5), Inches(0.9),
     "Find the real attack in millions of alerts — and explain why it is suspicious", size=26, color=WHITE)
text(s, Inches(0.8), Inches(3.9), Inches(11), Inches(1.2),
     "AIC × CrowdStrike Hackathon  ·  Track 2: Detect (AI-powered Blue Team)", size=18,
     color=RGBColor(0xB8, 0xC4, 0xD6))
for k, (big, small) in enumerate([("2.66M", "alerts in"), ("13", "incidents out"), ("8 / 8", "networks: #1 incident\nis the real attack")]):
    x = Inches(0.8 + 3.3 * k)
    text(s, x, Inches(5.1), Inches(3), Inches(0.8), big, size=40, color=RED, bold=True)
    text(s, x, Inches(5.9), Inches(3), Inches(0.8), small, size=15, color=WHITE)
s.notes_slide.notes_text_frame.text = (
    "AlertLens takes the millions of alerts a SOC receives, finds the few that are a real attack, groups them into "
    "incidents and explains, in plain language, why each one is suspicious.")

# 2 ─ problem
s = new_slide("Security teams drown in alerts", "The problem", """
A SOC (Security Operations Center) is the team that watches company systems for attacks. Its sensors raise far more
alerts than people can check. In our dataset: 8 company networks, 5 days, 2.66 million alerts. Only about 0.45% belong
to the real attack. The chart shows one network: the attack (red) does not stand out by volume; normal activity makes
bigger bursts every day. This is alert fatigue: real attacks get lost in the noise.""")
picture(s, "volume.png", Inches(0.5), Inches(1.75), Inches(8.6), Inches(4.9))
stat(s, Inches(9.5), Inches(1.9), Inches(3.5), "2.66M", "alerts from 3 sensors\nin 8 networks, 5 days")
stat(s, Inches(9.5), Inches(3.55), Inches(3.5), "0.45%", "of them are the\nreal attack")
stat(s, Inches(9.5), Inches(5.2), Inches(3.5), "12,017", "items to review today\n(sort by severity) to find 90%")

# 3 ─ severity
s = new_slide("Sorting by severity does not work", "Why the usual approach fails", """
What a SOC usually does first is sort alerts by severity. But attack and normal alerts have the same low levels,
mostly 3 and 6. The most dangerous attack steps, like switching user or running sudo, look like routine level-3 events.
So sorting by severity ranks the attack almost randomly: PR-AUC 0.006.""")
picture(s, "severity.png", Inches(0.5), Inches(1.75), Inches(8.2), Inches(4.9))
callout(s, Inches(9.1), Inches(2.0), Inches(3.7), Inches(1.6), "Attack steps like\n“web server became a user”\nare only severity 3")
callout(s, Inches(9.1), Inches(3.9), Inches(3.7), Inches(1.6), "We need context:\nis this unusual\nfor THIS network, NOW?")

# 4 ─ pipeline
s = new_slide("Our solution in one picture", "How AlertLens works", """
Walk left to right. 1: collect alerts from the three sensors into one table. 2: merge duplicates. 3: describe every
alert with about 100 answers to an analyst's questions. 4: a machine-learning model scores how likely each one is an
attack. 5: we explain every score. 6: suspicious alerts are chained into incidents. 7: a local AI writes the report.
The dashboard shows it all. The ML model detects; the AI explains.""")
steps = [("1", "Collect", "Wazuh, Suricata,\nAMiner → one table"), ("2", "Merge", "duplicates into\nmeta-alerts"),
         ("3", "Ask", "~100 analyst\nquestions per alert"), ("4", "Score", "ML model:\nhow likely attack?"),
         ("5", "Explain", "baseline vs. now\n+ reasons"), ("6", "Group", "alerts into\nattack incidents"),
         ("7", "Report", "local AI writes\nthe incident report")]
cw, gap, y0 = Inches(1.66), Inches(0.13), Inches(2.3)
for k, (n, head, sub) in enumerate(steps):
    x = Inches(0.6) + k * (cw + gap)
    hl = k == 3
    b = box(s, x, y0, cw, Inches(2.6), NAVY if hl else LIGHT)
    text(s, x, y0 + Inches(0.2), cw, Inches(0.6), n, size=30, color=RED, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, y0 + Inches(0.85), cw, Inches(0.5), head, size=21, color=WHITE if hl else NAVY, bold=True,
         align=PP_ALIGN.CENTER)
    text(s, x + Inches(0.08), y0 + Inches(1.45), cw - Inches(0.16), Inches(1), sub, size=13,
         color=WHITE if hl else GRAY, align=PP_ALIGN.CENTER)
    if k < len(steps) - 1:
        a = box(s, x + cw - Inches(0.02), y0 + Inches(1.15), gap + Inches(0.04), Inches(0.3), RED,
                shape=MSO_SHAPE.RIGHT_ARROW)
callout(s, Inches(0.6), Inches(5.35), Inches(12.1), Inches(1.0),
        "The machine-learning model decides what is suspicious.  The local AI explains it.  "
        "All data stays on the machine.", size=19)

# 5 ─ merge
s = new_slide("Step 1–2: from millions of alerts to a short list", "Merge duplicates", """
One attacker action can fire the same alert thousands of times. Example: the attacker's directory-guessing tool made
7,155 identical alerts in a few minutes; we merge them into one meta-alert (same network, host, rule and source,
less than 5 minutes apart). That takes us from 2.66 million to 99 thousand items. The model then flags 451, and those
are grouped into 13 incidents across 8 networks.""")
picture(s, "funnel.png", Inches(0.5), Inches(1.8), Inches(8.3), Inches(4.8))
callout(s, Inches(9.1), Inches(2.0), Inches(3.7), Inches(2.0),
        "Example: 7,155 alerts from the attacker’s website scanner → 1 meta-alert")
callout(s, Inches(9.1), Inches(4.3), Inches(3.7), Inches(2.0),
        "451 flagged, 91% are real attack.\nOnly 10 are truly outside the attack.")

# 6 ─ features
s = new_slide("Step 3: ask the questions an analyst would ask", "Features", """
The model never sees raw IPs or hostnames. Each meta-alert becomes about 100 numbers, each answering an analyst's
question. They're computed by small blocks called calcers, an idea we reused from the Magnit hackathon. Every answer
uses only alerts that happened before, so it works live. The charts show why this works: attack alerts come from rules
almost never seen here, at moments when many rare things happen at once, and their raw logs show attacker behaviour.""")
qs = ["Is this rule rare|in this network?", "What else unusual|is happening right now?",
      "Is the source IP new?|How many hosts has it touched?", "Do the raw logs show|attacker or admin behaviour?",
      "Is the user account new|on this host?"]
bullets(s, Inches(0.6), Inches(1.8), Inches(4.4), Inches(4.8), qs, size=17, gap=14)
picture(s, "features.png", Inches(5.0), Inches(1.8), Inches(7.9), Inches(4.3))
text(s, Inches(5.0), Inches(6.15), Inches(7.9), Inches(0.5), "blue = normal alerts,  red = attack alerts",
     size=13, color=GRAY, align=PP_ALIGN.CENTER)

# 7 ─ tree
s = new_slide("Step 4: the model learns which answers mean “attack”", "Training", """
Supervised learning: we show the computer examples from 7 networks together with the answer key (which alerts were the
attack), and it finds the questions that separate attack from normal by itself. This is a tiny 2-level tree it learned:
if many rare detections happen at once AND the rule is rarely active here, then attack. Our real model, LightGBM,
builds 800 small trees like this one after another; each new tree fixes the mistakes of the previous ones, and their
votes add up to the final score.""")
picture(s, "tree.png", Inches(0.4), Inches(1.75), Inches(8.4), Inches(4.9))
callout(s, Inches(9.1), Inches(1.9), Inches(3.7), Inches(1.45), "The computer found these questions itself, from examples")
callout(s, Inches(9.1), Inches(3.55), Inches(3.7), Inches(1.45), "Real model (LightGBM):\n800 trees, each fixing the last one’s mistakes")
callout(s, Inches(9.1), Inches(5.2), Inches(3.7), Inches(1.45), "Rare attacks get extra weight; 5 copies averaged")

# 8 ─ LOSO
s = new_slide("The honest test: a network the model has never seen", "Validation", """
We never test on data the model trained on. We hold out one whole network, train on the other 7, score the held-out
one, and repeat 8 times so every network is tested once. This matches real life: install AlertLens at a new company
and it must work from day one. Every number in this presentation comes from a network the model never saw.""")
picture(s, "loso.png", Inches(0.5), Inches(1.75), Inches(7.6), Inches(4.95))
bullets(s, Inches(8.5), Inches(2.1), Inches(4.4), Inches(4.5), [
    "Train|on 7 networks", "Test|on the 8th, which the model has never seen", "Repeat|8 times: every network is tested once",
    "No memorising|of IPs or hostnames. It has to generalise like a real deployment"], size=18, gap=16)

# 9 ─ results
s = new_slide("Results: 64 items to review instead of 12,017", "Results on unseen networks", """
PR-AUC measures how clean the top of our ranked list is: 1.0 means every real attack is above every false alarm.
We score 0.96. A generic anomaly detector (Isolation Forest) scores 0.37, and sorting by severity 0.006. To catch 90%
of the attack, an analyst reviews 64 items instead of 12,017. The top 20 is 100% real attack, and in all 8 networks
our #1 incident is the real attack.""")
picture(s, "results.png", Inches(0.4), Inches(1.75), Inches(12.5), Inches(3.9))
for k, (big, small) in enumerate([("0.96", "PR-AUC (1.0 = perfect)"), ("100%", "of the top 20 are real attack"),
                                  ("8 / 8", "networks: #1 incident is the attack"), ("90%", "attack stage predicted correctly")]):
    stat(s, Inches(0.8 + 3.1 * k), Inches(5.55), Inches(2.9), big, small)

# 10 ─ kill chain
s = new_slide("A real attack, step by step (santos, 17 January)", "Example", """
This is the real attack in the santos network, as AlertLens shows it. 11:21: the attacker scans the network and the
website. 11:23: uploads a backdoor (webshell) through WordPress. 11:24: cracks a user's password. 11:58: opens a remote
shell, and the web-server account switches to the real user gmorgan, then runs sudo as root. Each alert alone is low
severity. AlertLens links 11,218 raw alerts into one incident with 6 stages, ranked #1 in santos.""")
picture(s, "killchain.png", Inches(0.4), Inches(1.7), Inches(12.5), Inches(4.6))
callout(s, Inches(0.6), Inches(6.3), Inches(12.1), Inches(0.62),
        "Each alert alone looks minor (severity 3). Together they form a 6-stage attack chain → incident #1, confidence HIGH",
        size=16)

# 11 ─ baseline vs now
rows = json.load(open(FIGS / "baseline_table.json"))
s = new_slide("Why is it suspicious? Normal for this network vs. this alert", "Baseline · learning period · why", """
This answers the CrowdStrike judges' question. The first 24 hours of each network are its learning period: that's
how we learn what normal looks like (no attack starts before hour 46). Then for every alert we show what's normal here
next to this alert. Example: this rule is normally active in 71% of hours, but this one had never fired before; normally
0 rare detections happen at once, here 13. The last column is how much each difference pushed the score.""")
tbl = s.shapes.add_table(len(rows) + 1, 4, Inches(0.6), Inches(1.85), Inches(8.6), Inches(3.9)).table
for j, (h, wd) in enumerate([("Question", 4.0), ("Normal in santos", 1.8), ("This alert", 1.4), ("Push", 1.4)]):
    tbl.columns[j].width = Inches(wd)
    c = tbl.cell(0, j)
    c.text = h
    c.fill.solid()
    c.fill.fore_color.rgb = NAVY
for i, r in enumerate(rows, 1):
    for j, v in enumerate(r):
        c = tbl.cell(i, j)
        c.text = v
        c.fill.solid()
        c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
for i in range(len(rows) + 1):
    for j in range(4):
        p = tbl.cell(i, j).text_frame.paragraphs[0]
        for r in p.runs:
            r.font.size, r.font.name = Pt(15), FONT
            r.font.bold = i == 0 or j == 2
            r.font.color.rgb = WHITE if i == 0 else (RED if j in (2, 3) else NAVY)
callout(s, Inches(9.6), Inches(1.85), Inches(3.2), Inches(1.9),
        "Learning period:\nfirst 24 h of each network = what “normal” looks like")
callout(s, Inches(9.6), Inches(3.95), Inches(3.2), Inches(1.8),
        "After that, every feature uses only the past, so it works live")
text(s, Inches(0.6), Inches(5.95), Inches(12), Inches(0.5),
     "Alert: santos, 17 Jan 11:58:17. Raw log: “User successfully changed UID … www-data:gmorgan”  (the web server became a real user)",
     size=14, color=GRAY)

# 12 ─ SHAP receipt
s = new_slide("Every score comes with an itemised receipt", "Explainability (SHAP)", """
SHAP splits the model's score into the contribution of each reason, like an itemised receipt. It starts from a typical
alert, about 0.002% chance of attack. Then each reason pushes the score up (red) or down (blue). The rule had never
been active here: big push. 13 rare detections at the same time: big push. Attacker behaviour in the raw log: push.
Final score: 99%. The numbers add up exactly to the score the analyst sees.""")
picture(s, "shap_attack.png", Inches(0.4), Inches(1.7), Inches(8.9), Inches(5.0))
callout(s, Inches(9.6), Inches(1.9), Inches(3.2), Inches(1.5), "Starts at a typical alert:\n≈ 0.002% chance")
callout(s, Inches(9.6), Inches(3.6), Inches(3.2), Inches(1.5), "Each reason pushes\nup (red) or down (blue)")
callout(s, Inches(9.6), Inches(5.3), Inches(3.2), Inches(1.3), "Adds up exactly\nto the final 99%")

# 13 ─ rare but normal
s = new_slide("Rare is not the same as malicious", "Avoiding false alarms", """
The hard part is not flagging normal admin work. Left: an admin's SSH key login on a mail server. The rule is rare
there, so a "rarest rule first" detector would flag it, but nothing else is unusual, and the log shows an
administrator's key login, so our model scores it 0.1%. Right: the same user gmorgan ran sudo for the first time a day
earlier, from his home folder, to read a log. The model flags it (0.75), but it's a single-stage incident with no
attacker signs: confidence LOW, and the AI analyst calls it LIKELY BENIGN.""")
picture(s, "shap_admin.png", Inches(0.4), Inches(1.7), Inches(7.6), Inches(4.6))
box(s, Inches(8.4), Inches(1.85), Inches(4.4), Inches(4.75), LIGHT)
text(s, Inches(8.65), Inches(2.0), Inches(4), Inches(0.5), "Same user, same rule: first-time sudo", size=17, bold=True)
text(s, Inches(8.65), Inches(2.6), Inches(4), Inches(0.4), "ATTACK  ·  17 Jan", size=14, color=RED, bold=True)
text(s, Inches(8.65), Inches(2.95), Inches(4), Inches(1.0),
     "sudo from the website’s uploads folder, right after the web server became gmorgan → 6-stage incident, HIGH",
     size=14)
text(s, Inches(8.65), Inches(4.2), Inches(4), Inches(0.4), "NORMAL  ·  16 Jan", size=14, color=BLUE, bold=True)
text(s, Inches(8.65), Inches(4.55), Inches(4), Inches(1.2),
     "sudo from /home/gmorgan to read a log file → single-stage incident, LOW. AI analyst: “LIKELY BENIGN”", size=14)
text(s, Inches(0.6), Inches(6.35), Inches(7.4), Inches(0.5), "Only 10 flagged items in all 8 networks are truly outside the attack",
     size=15, color=NAVY, bold=True)

# 14 ─ LLM role
s = new_slide("The role of the AI (LLM): it explains, it does not detect", "Local LLM analyst", """
The ML model detects: it scores all 99 thousand meta-alerts in seconds, and its accuracy is measured. The LLM (Qwen3
14B, running locally with Ollama) receives only the top evidence of one incident and writes the report an analyst
would write: timeline, attacker IP, MITRE techniques, what to do next. The analyst can also chat with it. Why not let
the LLM read the alerts directly? At 1 second per alert, 2.66 million alerts would take about 30 days, and an LLM
can't see what's normal for a network from one alert. We keep it honest: it only uses our facts, it's told that rare
is not malicious, and it can only say TRUE POSITIVE for high-confidence incidents.""")
for k, (head, color, items) in enumerate([
        ("ML model  ·  DETECTS", NAVY, ["Scores all 99,161 meta-alerts|in seconds", "Learns what is normal|per network",
                                         "Accuracy is measured|on unseen networks"]),
        ("Local LLM  ·  EXPLAINS", RED, ["Writes the incident report:|timeline, IOCs, MITRE, next steps",
                                          "Answers the analyst’s questions|in a chat",
                                          "Runs on the laptop:|no data leaves the company"])]):
    x = Inches(0.6 + 6.2 * k)
    box(s, x, Inches(1.85), Inches(5.9), Inches(0.7), color)
    text(s, x, Inches(1.85), Inches(5.9), Inches(0.7), head, size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    box(s, x, Inches(2.55), Inches(5.9), Inches(2.45), LIGHT)
    bullets(s, x + Inches(0.3), Inches(2.75), Inches(5.4), Inches(2.2), items, size=17, gap=12)
callout(s, Inches(0.6), Inches(5.3), Inches(12.1), Inches(1.15),
        "Kept honest: grounded only in our evidence, told “rare ≠ malicious”.  "
        "Result: 9 / 9 real attack incidents → TRUE POSITIVE,  4 / 4 normal incidents → LIKELY BENIGN or NEEDS REVIEW",
        size=17)

# 15 ─ ablation
s = new_slide("What really matters: behaviour in this network", "Ablation", """
Ablation means removing one ingredient, retraining, and checking how much worse it gets. If the model only knows WHAT
a rule is (its name and category), PR-AUC drops to 0.49. Knowing how it behaves in this network is what gets us to 0.96.
Removing the time-window context costs only 0.02, so the model isn't just learning when the attack happened.""")
picture(s, "ablation.png", Inches(0.5), Inches(1.8), Inches(8.2), Inches(4.7))
callout(s, Inches(9.1), Inches(2.0), Inches(3.7), Inches(2.0), "Knowing only WHAT the rule is: 0.49")
callout(s, Inches(9.1), Inches(4.3), Inches(3.7), Inches(2.0), "Knowing HOW it behaves in this network: 0.96")

# 16 ─ limits
s = new_slide("Honest limits and next steps", "What we would do next", """
Be upfront: the labels are time windows, so some attack traffic is labeled normal; that's most of our false positives.
DNS exfiltration is mostly missed, because the detector that would catch it fires all day on normal traffic. The
high-confidence rule was chosen after seeing results. And it's 8 networks with one attack type: more variety is needed.""")
bullets(s, Inches(0.6), Inches(1.9), Inches(6.0), Inches(4.8), [
    "Labels are time windows:|some attack traffic is labeled “normal”", "DNS exfiltration is mostly missed:|its detector fires all day on normal traffic",
    "Confidence rule|was chosen after seeing results", "8 networks, one attack type:|needs more variety"], size=18, gap=16)
box(s, Inches(7.0), Inches(1.9), Inches(5.8), Inches(4.6), LIGHT)
text(s, Inches(7.3), Inches(2.05), Inches(5.3), Inches(0.5), "Next steps", size=20, bold=True, color=RED)
bullets(s, Inches(7.3), Inches(2.7), Inches(5.3), Inches(3.7), [
    "Per-domain features|to catch DNS exfiltration", "Analyst feedback|(“not an attack”) retrains the model",
    "Live streaming:|features already use only the past", "More attack types|and networks"], size=17, gap=14)

# 17 ─ closing
page += 1
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, H, NAVY, shape=MSO_SHAPE.RECTANGLE)
text(s, Inches(0.8), Inches(1.2), Inches(11.5), Inches(1.6),
     "From 2.66 million noisy alerts to a short ranked list of incidents, each with the reason why it is suspicious.",
     size=32, color=WHITE, bold=True)
box(s, Inches(0.8), Inches(3.05), Inches(1.4), Inches(0.09), RED, shape=MSO_SHAPE.RECTANGLE)
for k, (big, small) in enumerate([("64", "items to review\ninstead of 12,017"), ("0.96", "PR-AUC on networks\nthe model never saw"),
                                  ("100%", "local: no data\nleaves the company")]):
    x = Inches(0.8 + 3.8 * k)
    text(s, x, Inches(3.6), Inches(3.5), Inches(0.9), big, size=48, color=RED, bold=True)
    text(s, x, Inches(4.55), Inches(3.5), Inches(0.9), small, size=17, color=WHITE)
text(s, Inches(0.8), Inches(6.0), Inches(11), Inches(0.8), "Thank you  ·  Questions?  ·  Live demo: streamlit run app.py",
     size=22, color=RGBColor(0xB8, 0xC4, 0xD6))
s.notes_slide.notes_text_frame.text = "Switch to the dashboard for the live demo."

out = HERE / "AlertLens.pptx"
prs.save(out)
print("saved", out, "with", len(prs.slides), "slides")
