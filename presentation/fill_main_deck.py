"""Fill the team's Presentation.pptx template with AlertLens content. Run from the repo root."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

DECK = Path("Presentation.pptx")
FIGS = Path("presentation/figs")
WHITE = RGBColor(0xF4, 0xF4, 0xF4)

SLIDES = {
    "about": ("About Our Project",
              "Track 2 · Detect: an AI blue-team assistant that turns millions of security alerts into a few "
              "explained incidents.",
              ["Input: 2.66M alerts from Wazuh, Suricata and AMiner (AIT-ADS, 8 networks)",
               "Output: ranked incidents with reasons, MITRE stages and an analyst report",
               "Runs locally: no telemetry leaves the SOC"],
              "dashboard_incidents.png"),
    "problem": ("The Problem",
                "Security teams drown in alerts. The real attack is a tiny fraction of the noise.",
                ["2.66M alerts; only 450 of 99K alert groups are the attack",
                 "Analysts cannot read them all, so real attacks get missed (alert fatigue)",
                 "Hackers hide inside this noise"],
                "volume.png"),
    "severity": ("Severity Is Not Enough",
                 "Attack alerts and normal alerts get the same severity levels, so sorting by severity does not "
                 "find the attack.",
                 ["Wazuh severity ranks the attack worse than random (ROC-AUC 0.41)",
                  "By severity: about 12,017 items to review to catch 90% of the attack",
                  "We need context: what is normal for this network?"],
                 "severity.png"),
    "solution": ("Our Solution",
                 "AlertLens shrinks 2.66M alerts to 13 ranked incidents and explains each one.",
                 ["Merge duplicates into 99K meta-alerts",
                  "Machine learning flags 451; 91% of them are the real attack",
                  "Chain them into 13 incidents: #1 is the real attack in 8 of 8 networks",
                  "Threat hunt: unsupervised hunting over raw logs"],
                 "funnel.png"),
    "how": ("How It Works",
            "The model learns what attacker behaviour looks like from context features, computed from past alerts "
            "only.",
            ["Rarity and novelty: first time a rule, IP or account is seen",
             "Bursts and several sensors firing within 10 minutes",
             "Attacker vs. admin hints in raw logs (www-data, /etc/shadow, scanners)",
             "LightGBM, 5-seed ensemble"],
            "features.png"),
    "loso": ("Tested On Unseen Networks",
             "We train on 7 networks and test on the 8th, rotating through all 8 (leave-one-network-out).",
             ["Every score comes from a model that never saw that network",
              "Like deploying to a brand-new customer",
              "No peeking at the future: features use only earlier alerts"],
             "loso.png"),
    "explain": ("Explainable Scores",
                "Every alert gets plain-English reasons (SHAP): which signals pushed its score up or down.",
                ["Example: www-data switching to user gmorgan, a first-time event, is scored as attack",
                 "Analysts see why, not just a number",
                 "Mark a false positive once and it is suppressed on that host"],
                "shap_attack.png"),
    "story": ("The Attack Story",
              "Flagged alerts are chained into an incident timeline with kill-chain stages and MITRE ATT&CK "
              "techniques.",
              ["From reconnaissance to webshell upload to privilege escalation",
               "A local LLM (qwen3:14b via Ollama) writes the report and answers questions",
               "Confidence HIGH only with high risk and 2+ attack stages"],
              "killchain.png"),
    "results": ("Results and Evaluation",
                "Averaged over 8 unseen networks, AlertLens beats common approaches by a wide margin.",
                ["PR-AUC 0.96 vs. 0.37 Isolation Forest, 0.14 rarity, 0.01 Wazuh severity",
                 "Top 20 alerts: 100% are real attack",
                 "Catch 90% of the attack: 64 items to review instead of 12,017",
                 "Attack stage predicted correctly 90% of the time"],
                "results.png"),
    "impact": ("Real World Impact",
               "Analysts review about 64 items instead of 12,000 and still catch 90% of the attack.",
               ["Hides 99.91% of benign meta-alerts at 90% attack recall",
                "Normal admin work is explained as benign, not escalated",
                "Private: model and LLM run on-premises, no data leaves the SOC"],
               "shap_admin.png"),
    "wins": ("Why Our Product Wins",
             "Accurate, explainable, tested on unseen networks, and private.",
             ["Behaviour matters: PR-AUC 0.49 from rule meaning alone, 0.96 with context",
              "Every score has reasons, every incident has a story",
              "End to end: raw logs to analyst report in one dashboard"],
             "ablation.png"),
}
EXISTING = {1: "about", 2: "problem", 3: "solution", 4: "how", 6: "results", 7: "impact", 8: "wins"}
NEW_AFTER = {"problem": ["severity"], "how": ["loso", "explain", "story"]}


def placeholders(slide):
    return {sh.placeholder_format.type: sh for sh in slide.placeholders}


def fill(slide, key):
    heading, sub, bullets, fig = SLIDES[key]
    ph = placeholders(slide)
    if PP_PLACEHOLDER.PICTURE in ph:
        el = ph[PP_PLACEHOLDER.PICTURE]._element
        el.getparent().remove(el)
    if PP_PLACEHOLDER.SUBTITLE in ph:
        ph[PP_PLACEHOLDER.SUBTITLE].text_frame.text = "ALERTLENS · AIC × CROWDSTRIKE · TRACK 2"

    title = ph[PP_PLACEHOLDER.TITLE]
    title.left, title.top, title.width, title.height = Inches(0.13), Inches(0.45), Inches(9.7), Inches(1.05)
    tf = title.text_frame
    tf.clear()
    tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.TOP
    tf.paragraphs[0].text = heading
    tf.paragraphs[0].runs[0].font.size = Pt(26)
    p = tf.add_paragraph()
    p.text = sub
    p.runs[0].font.size = Pt(13)

    if PP_PLACEHOLDER.BODY in ph:
        el = ph[PP_PLACEHOLDER.BODY]._element
        el.getparent().remove(el)
    body = slide.shapes.add_textbox(Inches(6.7), Inches(1.65), Inches(3.15), Inches(3.75))
    bf = body.text_frame
    bf.word_wrap, bf.vertical_anchor = True, MSO_ANCHOR.TOP
    for i, b in enumerate(bullets):
        p = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
        p.text = "• " + b
        p.runs[0].font.size = Pt(12)
        p.runs[0].font.color.rgb = WHITE
        p.space_after = Pt(9)

    w, h = Image.open(FIGS / fig).size
    box_l, box_t, box_w, box_h = 0.15, 1.65, 6.35, 3.75
    scale = min(box_w / w, box_h / h)
    pw, phh = w * scale, h * scale
    slide.shapes.add_picture(str(FIGS / fig), Inches(box_l + (box_w - pw) / 2), Inches(box_t + (box_h - phh) / 2),
                             Inches(pw), Inches(phh))
    slide.notes_slide.notes_text_frame.text = sub + "\n" + "\n".join(bullets)


def main():
    prs = Presentation(DECK)
    slides = list(prs.slides)
    layout = slides[1].slide_layout

    for sh in slides[0].placeholders:
        if not sh.has_text_frame:
            continue
        paras = sh.text_frame.paragraphs
        txt = sh.text_frame.text
        if txt.startswith("HACKATHON"):
            paras[0].runs[0].text, paras[1].runs[0].text = "ALERTLENS", "ALERT TRIAGE"
        elif txt.startswith("MM.DD.YY"):
            for para, t in zip(paras, ["10.12.26", "AIC × CROWDSTRIKE", "HACKATHON"]):
                if para.runs:
                    para.runs[0].text = t
                    for r in para.runs[1:]:
                        r.text = ""
        elif txt.startswith("PRESENTATOR"):
            for para, t in zip(paras, ["TRACK 2 · DETECT", "AI-POWERED BLUE TEAM"]):
                if para.runs:
                    para.runs[0].text = t
                    for r in para.runs[1:]:
                        r.text = ""

    for idx, key in EXISTING.items():
        fill(slides[idx], key)
    video = placeholders(slides[5])
    video[PP_PLACEHOLDER.SUBTITLE].text_frame.paragraphs[0].runs[0].text = "ALERTLENS · AIC × CROWDSTRIKE · TRACK 2"

    ids = prs.slides._sldIdLst
    order = list(ids)
    new_entries = {}
    for anchor, keys in NEW_AFTER.items():
        new_entries[anchor] = []
        for key in keys:
            s = prs.slides.add_slide(layout)
            fill(s, key)
            new_entries[anchor].append(ids[-1])
    for el in [e for v in new_entries.values() for e in v]:
        ids.remove(el)
    final = []
    for idx, el in enumerate(order):
        final.append(el)
        key = EXISTING.get(idx)
        final += new_entries.get(key, [])
    for el in list(ids):
        ids.remove(el)
    for el in final:
        ids.append(el)
    prs.save(DECK)
    print(f"saved {DECK} with {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
